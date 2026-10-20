"""HIST-03 parser-reproducibility check: re-parse the SAME workbook bytes with openpyxl.

This is evidence that two parsers read identical bytes identically. It is NOT a
second data source and proves nothing about the correctness of the World Bank
series itself. Output is metadata only (no series values).
"""
import io
from decimal import Decimal

from tools.hist_03_worldbank_validate import (END, SHEET, START, UNIT_RE, _period, canonical_sha256,
                                              expected_months, parse)

FIELDS = ("count_in_range", "first_in_range", "last_in_range", "gold_column", "header_text",
          "unit_accepted", "sheet_period_count", "sheet_first_period", "sheet_last_period",
          "canonical_series_sha256", "non_integer_count_in_range")


def parse_openpyxl(raw, start=START, end=END):
    """Independent parse via openpyxl; same output shape as validator.parse() (minus diagnostics)."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError("SHEET_NOT_FOUND")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    header = next((i for i, r in enumerate(rows) if any(v is not None and str(v).strip().upper() == "GOLD" for v in r)), None)
    if header is None:
        raise ValueError("GOLD_COLUMN_NOT_FOUND")
    idx = [i for i, v in enumerate(rows[header]) if v is not None and str(v).strip().upper() == "GOLD"]
    if len(idx) != 1:
        raise ValueError("AMBIGUOUS_GOLD_COLUMN")
    gi = idx[0]
    series, periods, unit_texts, seen = {}, [], [], False
    for r in rows[header + 1:]:
        d = "" if r[0] is None else str(r[0]).strip()
        raw_v = r[gi] if gi < len(r) else None
        period = _period(d) if d else None
        if period is None:
            if not seen and raw_v is not None and str(raw_v).strip():
                unit_texts.append(str(raw_v).strip())
            continue
        seen = True
        periods.append(period)
        if start <= period <= end and raw_v is not None and not isinstance(raw_v, bool):
            try:
                v = Decimal(repr(raw_v)) if isinstance(raw_v, float) else Decimal(str(raw_v).strip())
            except Exception:
                continue
            if v.is_finite() and v > 0:
                series[period] = v
    return {"gold_column": get_column_letter(gi + 1), "header_text": str(rows[header][gi]),
            "unit_texts": unit_texts, "series": series, "all_periods": periods}


def fingerprint(p, start=START, end=END):
    s = p["series"]
    months = [m for m in expected_months(start, end) if m in s]
    ap = p["all_periods"]
    return {
        "count_in_range": len(s), "first_in_range": min(s) if s else None, "last_in_range": max(s) if s else None,
        "gold_column": p["gold_column"], "header_text": p["header_text"],
        "unit_accepted": next((u for u in p["unit_texts"] if UNIT_RE.search(u)), None),
        "sheet_period_count": len(ap), "sheet_first_period": ap[0] if ap else None,
        "sheet_last_period": ap[-1] if ap else None,
        "canonical_series_sha256": canonical_sha256(s, months),
        "non_integer_count_in_range": sum(1 for v in s.values() if v != v.to_integral_value()),
    }


def compare(a, b):
    """Return the FIELDS on which two fingerprints differ (empty list means MATCH)."""
    return [f for f in FIELDS if a.get(f) != b.get(f)]


def crosscheck(raw, start=START, end=END):
    try:
        import openpyxl
        version = openpyxl.__version__
        other = fingerprint(parse_openpyxl(raw, start, end), start, end)
    except Exception as exc:  # parser unavailable or failed: never silently pass
        return {"status": "UNAVAILABLE_OR_FAILED", "error_class": type(exc).__name__, "error": str(exc)[:200]}
    primary = fingerprint(parse(raw, start, end), start, end)
    diff = compare(primary, other)
    return {"status": "MATCH" if not diff else "MISMATCH", "parser": "openpyxl " + version,
            "differing_fields": diff, "compared_fields": list(FIELDS),
            "independent_fingerprint": other,
            "note": "parser reproducibility on identical bytes; NOT a second data source"}
